from datetime import datetime
import os
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def track_study_metrics_in_excel(
    study_id: str,
    design_total_rows: int,
    design_total_forms: int,
    design_total_vars: int,
    unique_mapped_vars: int,
    total_output_rows: int,
    unmapped_output_rows: int,
    mapped_pct: float,
    output_file: str,
    excel_path: str = 'study_mapping_metrics.xlsx',
):
    """Tracks mapping metrics per study_id in Excel.

    If run multiple times for the same study_id, keeps ONLY the run with the
    highest match percentage (mapped_pct).
    """

    headers = [
        'Study ID',
        'Last Updated',
        'Design Spec Total Rows',
        'Total Unique Forms',
        'Total Unique Variables',
        'Mapped Variables',
        'Unmapped Output Rows',
        'Total Output Rows',
        'Match Percentage (%)',
        'Output File Path',
    ]

    # Normalize mapped_pct if passed as 0-100 percentage instead of decimal (0.0-1.0)
    pct_val = mapped_pct / 100.0 if mapped_pct > 1.0 else mapped_pct
    
    # NEW: Round to 4 decimal places (e.g., 0.672256 -> 0.6723). 
    # This ensures Excel displays exactly 67.23% without trailing numbers in the formula bar.
    pct_val = round(pct_val, 4)

    # Create or Load Workbook
    if not os.path.exists(excel_path):
        wb = openpyxl.Workbook()
        ws_tracker = wb.active
        ws_tracker.title = 'Study Metrics Tracker'
        ws_tracker.views.sheetView[0].showGridLines = True

        # Header Row
        ws_tracker.append([])  # Row 1 empty
        ws_tracker.append(headers)  # Row 2 Headers
    else:
        wb = openpyxl.load_workbook(excel_path)
        if 'Study Metrics Tracker' in wb.sheetnames:
            ws_tracker = wb['Study Metrics Tracker']
        else:
            ws_tracker = wb.create_sheet(title='Study Metrics Tracker')
            ws_tracker.append([])
            ws_tracker.append(headers)

    # Search for existing Study ID
    target_row = None
    existing_max_pct = -1.0

    for r in range(3, ws_tracker.max_row + 1):
        cell_study_id = ws_tracker.cell(row=r, column=1).value
        if cell_study_id and str(cell_study_id).strip() == str(study_id).strip():
            target_row = r
            raw_pct = ws_tracker.cell(row=r, column=9).value
            try:
                # NEW: Added safety check in case the cell was saved as a literal string like "67.23%"
                if isinstance(raw_pct, str) and '%' in raw_pct:
                    existing_max_pct = float(raw_pct.replace('%', '').strip()) / 100.0
                else:
                    existing_max_pct = float(raw_pct) if raw_pct is not None else -1.0
            except ValueError:
                existing_max_pct = -1.0
            break

    # Formatted as DD-MMM-YYYY HH:MM (e.g., 01-Jan-2026 14:30)
    timestamp = datetime.now().strftime('%d-%b-%Y %H:%M')

    new_row_data = [
        study_id,
        timestamp,
        design_total_rows,
        design_total_forms,
        design_total_vars,
        unique_mapped_vars,
        unmapped_output_rows,
        total_output_rows,
        pct_val,  # Now perfectly rounded to 4 decimals
        output_file,
    ]

    # Upsert Logic: High-match percentage retention
    if target_row is not None:
        if pct_val > existing_max_pct:
            print(
                f" Found higher match percentage for '{study_id}' ({pct_val*100:.2f}% > {existing_max_pct*100:.2f}%). Updating Excel record..."
            )
            for c_idx, val in enumerate(new_row_data, 1):
                ws_tracker.cell(row=target_row, column=c_idx, value=val)
        else:
            print(
                f" Existing record for '{study_id}' has a higher or equal match percentage ({existing_max_pct*100:.2f}% >= {pct_val*100:.2f}%). Excel tracker not overwritten."
            )
            wb.save(excel_path)
            return
    else:
        print(
            f" New Study ID '{study_id}' detected. Adding new record to Excel tracker..."
        )
        ws_tracker.append(new_row_data)
        target_row = ws_tracker.max_row

    # Styling & Formatting Setup
    HEADER_FILL = PatternFill(start_color='ADD8E6', end_color='ADD8E6', fill_type='solid')
    HEADER_FONT = Font(name='Segoe UI', size=11, bold=True, color='000000')
    ZEBRA_FILL = PatternFill(start_color='F9FAFB', end_color='F9FAFB', fill_type='solid')
    WHITE_FILL = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')
    THIN_BORDER = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9'),
    )

    # Format Headers
    ws_tracker.row_dimensions[2].height = 28
    for col_num in range(1, len(headers) + 1):
        cell = ws_tracker.cell(row=2, column=col_num)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    # Format Data Rows
    for r_idx in range(3, ws_tracker.max_row + 1):
        ws_tracker.row_dimensions[r_idx].height = 20
        fill = ZEBRA_FILL if r_idx % 2 == 0 else WHITE_FILL
        for c_idx in range(1, len(headers) + 1):
            cell = ws_tracker.cell(row=r_idx, column=c_idx)
            cell.fill = fill
            cell.border = THIN_BORDER
            cell.font = Font(name='Segoe UI', size=10)

            if c_idx in [1, 2]:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            elif c_idx in range(3, 9):
                cell.number_format = '#,##0'
                cell.alignment = Alignment(horizontal='right', vertical='center')
            elif c_idx == 9:
                cell.number_format = '0.00%'
                cell.alignment = Alignment(horizontal='right', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center')

    # Adjust Column Widths
    for col in ws_tracker.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row == 1:
                continue
            val_str = str(cell.value or '')
            if cell.number_format and '%' in cell.number_format:
                try:
                    val_str = f'{float(cell.value or 0)*100:.2f}%'
                except:
                    pass
            max_len = max(max_len, len(val_str))
        ws_tracker.column_dimensions[col_letter].width = max(max_len + 4, 14)
    
    ws_tracker.column_dimensions['J'].width = 35

    wb.save(excel_path)
    print(f' Excel tracker successfully saved to: {excel_path}\n')