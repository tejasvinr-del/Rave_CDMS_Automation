import re
import copy
import openpyxl
from datetime import datetime
from Study_track import track_study_metrics_in_excel
import pandas as pd
import numpy as np
from extarct_design import extract_rave_xml_to_df
from transpose_logic import replace_trans
from openpyxl.styles import Font, PatternFill, Border, Side

def generate_cdms_mapping():
    # ---------------------------------------------------------
    # CONFIGURATION & CONSTANTS
    # ---------------------------------------------------------
    design_spec_file = 'M26-167_M26-167_DRAFT_V0.2 2.xls'
    design_spec_file = 'ALS/' + design_spec_file
    design_spec_sheet = 'Fields'
    
    cdms_lib_file = 'Global_RAVE_CDMS_Library.xlsx'
    cdms_lib_sheet = 'Database Mapping'
    
    # Define your 3 reference files here (Set to None if not using them)
    ref_cdms_file1 = 'M26-266 Clinical Database Mapping Specification-V1.0 21AUG2026 (1).xlsx'
    ref_cdms_file2 = 'CDMS_preconf_golive/Oncology_CDMS/M26-120_Clinical Database Mapping Specification-V120_31Aug2026.xlsx'
    ref_cdms_file3 = 'CDMS_preconf_golive/Oncology_CDMS/M26-118 Clinical Database Mapping Specification-V1.0 20AUG2026.xlsx'
    # ref_cdms_sheet = 'Database Mapping'
#Aesthetics
    # ref_cdms_file1 = 'CDMS_preconf_golive/Aesthetics_CDMS/M26-479 Clinical Database Mapping Specification - 28 Jul 2026 V1.0.xlsx'
    # ref_cdms_file2 = 'CDMS_preconf_golive/Aesthetics_CDMS/M26-109 Clinical Database Mapping Specification - 31 Jul 2026 V1.0.xlsx'
    # ref_cdms_file3 = 'CDMS_preconf_golive/Aesthetics_CDMS/M23-066 Clinical Database Mapping Specification - 03 Aug 2026 V1.0.xlsx'
    # ref_cdms_file3 = 'CDMS_preconf_golive/Aesthetics_CDMS/M26-266 Clinical Database Mapping Specification-V1.0 21AUG2026.xlsx'

    ref_cdms_sheet = 'Database Mapping'

    studynum = design_spec_file.split('_')[0]
    studynum = studynum.split('/')[-1]
    current_datetime = datetime.now().strftime('%d %b %Y')
    output_file = f"CDMS_Generated/{studynum} Clinical Database Mapping Specification - {current_datetime} V1.0.xlsx"
    # output_file= f"final_output.xlsx"
    
    output_headers = [
        'Source System', 'Source Table', 'Source Column', 'Target Table', 
        'Target Column', 'DataType', 'Length', 'Precision', 
        'Mapping Logic (pseudo code)', 'Transpose', 'Additional details'
    ]
    
    mandatory_vars_list = [
        'DCMDATE', 'DOMAIN', 'GRPID', 'SPID', 'STUDYID', 
        'SUBJID', 'VISIT', 'VISITNUM', 'INVID'
    ]

    conditional_vars_list = [
        'TUEVAL','PRTRT','ECTRT''PRSTTPT','PRSTRTPT',
        'CMSTTPT','CMSTRTPT','CMTRT'
    ]

    print("Loading data files...")
    
    # ---------------------------------------------------------
    # STEP 1: LOAD INPUTS (WITH CUSTOM XML PARSER)
    # ---------------------------------------------------------
    try:
        cdms_df = pd.read_excel(cdms_lib_file, sheet_name=cdms_lib_sheet)
        approval_df = pd.read_excel(cdms_lib_file, sheet_name='Approval', header=None)
        doc_history_df = pd.read_excel(cdms_lib_file, sheet_name='Document History', header=None)
    except Exception as e:
        print(f"Error loading CDMS Library: {e}")
        return
        
    # Helper function to load reference files
    def load_ref_cdms(filepath):
        # Added check: If no filepath provided, return empty dataframe immediately
        if not filepath:
            return pd.DataFrame()
            
        try:
            df = pd.read_excel(filepath, sheet_name=ref_cdms_sheet)
            if 'Target Variable' in df.columns and 'Target Column' not in df.columns:
                df = df.rename(columns={'Target Variable': 'Target Column'})
            for col in output_headers:
                if col not in df.columns:
                    df[col] = None
                    
            df['Source Table'] = df['Source Table'].astype(str).str.strip()
            df['Source Column'] = df['Source Column'].astype(str).str.strip()
            df['Base_Source_Column'] = df['Source Column'].astype(str).str.replace(r'(_STD|_RAW|_U)$', '', regex=True)
            return df
        except Exception as e:
            print(f"Warning: Could not load Ref CDMS Library {filepath}: {e}")
            return pd.DataFrame()

    ref_cdms_df1 = load_ref_cdms(ref_cdms_file1)
    ref_cdms_df2 = load_ref_cdms(ref_cdms_file2)
    ref_cdms_df3 = load_ref_cdms(ref_cdms_file3)

    try:
        design_df = pd.read_excel(design_spec_file, sheet_name=design_spec_sheet)
    except ValueError as e:
        if "Excel file format cannot be determined" in str(e):
            print("\nDetected standard Rave XML export. Engaging custom XML parser...")
            try:
                design_df = extract_rave_xml_to_df(design_spec_file, design_spec_sheet)
                print("Successfully mapped XML data into Pandas DataFrame!\n")
            except Exception as xml_e:
                print(f"Custom XML parsing failed: {xml_e}")
                return
        else:
            print(f"Error loading Design Spec: {e}")
            return
    except Exception as e:
        print(f"Error loading Design Spec: {e}")
        return

    design_df = design_df.loc[:, design_df.columns.notna()]

    # 2. Filter for DraftFieldActive == TRUE
    if 'DraftFieldActive' in design_df.columns:
    # We cast to string and uppercase to safely catch 'TRUE', 'True', 'true', or boolean True
        mask = design_df['DraftFieldActive'].astype(str).str.upper() == 'TRUE'
        design_df = design_df[mask]
    
    # Optional: Reset the index after filtering so your row numbers remain sequential
        design_df = design_df.reset_index(drop=True)
    else:
        print("Warning: 'DraftFieldActive' column not found in the parsed sheet.")

    if 'Target Variable' in cdms_df.columns and 'Target Column' not in cdms_df.columns:
        cdms_df = cdms_df.rename(columns={'Target Variable': 'Target Column'})

    for col in output_headers:
        if col not in cdms_df.columns:
            cdms_df[col] = None

    cdms_df['Source Table'] = cdms_df['Source Table'].astype(str).str.strip()
    cdms_df['Source Column'] = cdms_df['Source Column'].astype(str).str.strip()

    # ---------------------------------------------------------
    # STEP 2: DERIVE SOURCE VARIABLES (Design Spec)
    # ---------------------------------------------------------
    if 'FormOID' in design_df.columns and 'FieldOID' in design_df.columns:
        design_df['Derived_Source_Table'] = (design_df['FormOID'].astype(str) + '_BIOSTAT').str.strip()
        design_df['Derived_Source_Column'] = design_df['FieldOID'].astype(str).str.strip()
        
        base_form = design_df['FormOID'].astype(str).str.replace(r'_?\d+(_F)$', r'\1', regex=True)
        base_form = base_form.str.replace(r'_?\d+$', '', regex=True)
        design_df['Base_Source_Table'] = base_form + '_BIOSTAT'
    else:
        print(f"Error: Could not find 'FormOID' or 'FieldOID' in the parsed data.")
        return

    # ---------------------------------------------------------
    # STEP 2b: EXPAND CDMS & REF LIBRARY TO ACCOMMODATE NUMBERED FORMS
    # ---------------------------------------------------------
    # Expand Primary CDMS Library
    lib_base_form = cdms_df['Source Table'].astype(str).str.replace(r'_BIOSTAT$', '', regex=True)
    lib_base_form = lib_base_form.str.replace(r'_?\d+(_F)$', r'\1', regex=True)
    lib_base_form = lib_base_form.str.replace(r'_?\d+$', '', regex=True)
    cdms_df['Base_Source_Table'] = lib_base_form + '_BIOSTAT'

    form_mapping = design_df[['Base_Source_Table', 'Derived_Source_Table']].drop_duplicates()
    forms_to_expand = form_mapping[form_mapping['Base_Source_Table'] != form_mapping['Derived_Source_Table']]
    
    expanded_cdms_rows = []
    for _, row in forms_to_expand.iterrows():
        base = row['Base_Source_Table']
        derived = row['Derived_Source_Table']
        base_rows = cdms_df[cdms_df['Base_Source_Table'] == base].copy()
        
        if not base_rows.empty:
            if not (cdms_df['Source Table'] == derived).any():
                base_rows['Source Table'] = derived
                expanded_cdms_rows.append(base_rows)
            
    if expanded_cdms_rows:
        cdms_df = pd.concat([cdms_df] + expanded_cdms_rows, ignore_index=True)

    # Helper function to expand reference libraries
    def expand_ref_library(ref_df):
        if ref_df.empty:
            return ref_df
        ref_base_form = ref_df['Source Table'].astype(str).str.replace(r'_BIOSTAT$', '', regex=True)
        ref_base_form = ref_base_form.str.replace(r'_?\d+(_F)$', r'\1', regex=True)
        ref_base_form = ref_base_form.str.replace(r'_?\d+$', '', regex=True)
        ref_df['Base_Source_Table'] = ref_base_form + '_BIOSTAT'
        
        expanded_ref_rows = []
        for _, row in forms_to_expand.iterrows():
            base = row['Base_Source_Table']
            derived = row['Derived_Source_Table']
            base_rows = ref_df[ref_df['Base_Source_Table'] == base].copy()
            
            if not base_rows.empty:
                if not (ref_df['Source Table'] == derived).any():
                    base_rows['Source Table'] = derived
                    expanded_ref_rows.append(base_rows)
                    
        if expanded_ref_rows:
            ref_df = pd.concat([ref_df] + expanded_ref_rows, ignore_index=True)
        return ref_df

    ref_cdms_df1 = expand_ref_library(ref_cdms_df1)
    ref_cdms_df2 = expand_ref_library(ref_cdms_df2)
    ref_cdms_df3 = expand_ref_library(ref_cdms_df3)


    # ---------------------------------------------------------
    # STEP 3 & 4: MATCH WITH PRIMARY CDMS & CASCADE THROUGH REF FILES
    # ---------------------------------------------------------
    print("Joining Design Spec with Primary CDMS Library...")
    cdms_df['Base_Source_Column'] = cdms_df['Source Column'].astype(str).str.replace(r'(_STD|_RAW|_U)$', '', regex=True)
    
    mapped_temp_df = pd.merge(
        design_df,
        cdms_df,
        left_on=['Derived_Source_Table', 'Derived_Source_Column'],
        right_on=['Source Table', 'Base_Source_Column'],
        how='inner'
    )
    
    core_mapped_df = mapped_temp_df[output_headers].copy()
    
    mapped_keys = mapped_temp_df[['Derived_Source_Table', 'Derived_Source_Column']].drop_duplicates()
    merged_for_unmatched = pd.merge(
        design_df, 
        mapped_keys, 
        on=['Derived_Source_Table', 'Derived_Source_Column'], 
        how='left', 
        indicator=True
    )
    unmatched_design = merged_for_unmatched[merged_for_unmatched['_merge'] == 'left_only'].copy()
    unmatched_design = unmatched_design.drop(columns=['_merge'])
    
    # 4. Sequential Check through Reference CDMS files
    ref_files_list = [(1, ref_cdms_df1), (2, ref_cdms_df2), (3, ref_cdms_df3)]
    
    for i, ref_df in ref_files_list:
        if not ref_df.empty and not unmatched_design.empty:
            print(f"Checking Reference CDMS {i} for remaining unmapped variables...")
            
            ref_mapped_temp = pd.merge(
                unmatched_design,
                ref_df,
                left_on=['Derived_Source_Table', 'Derived_Source_Column'],
                right_on=['Source Table', 'Base_Source_Column'],
                how='inner'
            )
            
            if not ref_mapped_temp.empty:
                print(f" -> Recovered {len(ref_mapped_temp)} matches from Reference CDMS {i} - {eval(f'ref_cdms_file{i}')}.")
                core_mapped_df = pd.concat([core_mapped_df, ref_mapped_temp[output_headers]], ignore_index=True)
                
                # Update unmatched design df by filtering out what we just mapped
                ref_mapped_keys = ref_mapped_temp[['Derived_Source_Table', 'Derived_Source_Column']].drop_duplicates()
                ref_mapped_keys['found_in_ref'] = True
                
                unmatched_design = pd.merge(
                    unmatched_design,
                    ref_mapped_keys,
                    on=['Derived_Source_Table', 'Derived_Source_Column'],
                    how='left'
                )
                unmatched_design = unmatched_design[unmatched_design['found_in_ref'].isna()].copy()
                unmatched_design = unmatched_design.drop(columns=['found_in_ref'])
    
    unmatched_df = pd.DataFrame(columns=output_headers)
    if not unmatched_design.empty:
        unmatched_df['Source Table'] = unmatched_design['Derived_Source_Table']
        unmatched_df['Source Column'] = unmatched_design['Derived_Source_Column']
    unmatched_df['is_unmatched'] = True

    # ---------------------------------------------------------
    # STEP 5, 6, 7 & 8: BUILD STRUCTURAL DOMAINS BEFORE CONDITIONAL RULES
    # ---------------------------------------------------------
    # 5. SDQ Rule
    sdq_df = cdms_df[cdms_df['Source Table'] == 'SDQ'].copy()
    if not sdq_df.empty: core_mapped_df = pd.concat([core_mapped_df, sdq_df[output_headers]], ignore_index=True)

    # 6. External Rule
    all_external_domains = ['DA','SP','RN','RNKIT','LAB','OBSERVATIONS','PK_DATA']
    external_df = cdms_df[cdms_df['Source Table'].isin(all_external_domains)].copy()
    if not external_df.empty:
        core_mapped_df = pd.concat([core_mapped_df, external_df[output_headers]], ignore_index=True)
        ext_target_tables = external_df['Target Table'].dropna().unique()
        ext_derived_df = cdms_df[(cdms_df['Source Table'] == 'DERIVED') & (cdms_df['Target Table'].isin(ext_target_tables))].copy()
        if not ext_derived_df.empty: core_mapped_df = pd.concat([core_mapped_df, ext_derived_df[output_headers]], ignore_index=True)

    # 7. Default Domains Rule
    all_default_domains = ['SYSDT', 'SV', 'SUBJECT_VISIT_BYDOMAIN']
    default_domains_df = cdms_df[cdms_df['Target Table'].isin(all_default_domains)].copy()
    if not default_domains_df.empty: core_mapped_df = pd.concat([core_mapped_df, default_domains_df[output_headers]], ignore_index=True)

    # 8. LB Cascading Rule
    lb_tables = ['LB_SAE', 'LB', 'LB_LOCAL', 'QS_ECOA', 'QS_ECOA_AA', 'QS_ECOA_HS']
    mapped_lb_vars = core_mapped_df[core_mapped_df['Target Table'].isin(lb_tables)][['Target Table', 'Target Column']].dropna().drop_duplicates()
    
    if not mapped_lb_vars.empty:
        mapped_lb_vars = mapped_lb_vars.rename(columns={'Target Table': 'Source Table', 'Target Column': 'Source Column'})
        matched_lb_cascade_df = pd.merge(mapped_lb_vars, cdms_df, on=['Source Table', 'Source Column'], how='inner')
        final_lb_cascade_df = matched_lb_cascade_df[output_headers].drop_duplicates()
        core_mapped_df = pd.concat([core_mapped_df, final_lb_cascade_df], ignore_index=True)

    # ---------------------------------------------------------
    # CONDITIONAL RULES (Requires all tables established)
    # ---------------------------------------------------------
    appended_rows_list = []
    active_table_pairs = core_mapped_df[['Source Table', 'Target Table']].dropna().drop_duplicates()
    unique_target_tables = core_mapped_df['Target Table'].dropna().unique()

    # FIX: Create a Combined Library for Conditional matching that includes the Primary and all 3 Ref files
    dfs_to_concat = [cdms_df]
    for _, ref_df in ref_files_list:
        if not ref_df.empty:
            dfs_to_concat.append(ref_df)
            
    combined_lib_df = pd.concat(dfs_to_concat, ignore_index=True).drop_duplicates(subset=['Source Table', 'Target Table', 'Target Column'])

    # Pre-calculate masks against the Combined Library
    is_mand_list = combined_lib_df['Target Column'].isin(mandatory_vars_list)
    is_missing_src_col = (
        combined_lib_df['Source Column'].isna() | 
        (combined_lib_df['Source Column'].astype(str).str.strip() == '') | 
        (combined_lib_df['Source Column'].astype(str).str.lower() == 'nan')
    )

    is_xx_spid = combined_lib_df['Target Column'].astype(str).str.match(r'^[A-Za-z]{2,3}(SPID)$') 
    is_xx_cat_scat = (
        combined_lib_df['Target Column'].astype(str).str.match(r'^[A-Za-z]{2,3}(CAT|SCAT)$') & 
        (combined_lib_df['Source Column'].isna() | combined_lib_df['Source Column'].astype(str).str.strip().isin(['', 'DATAPAGENAME']))
    )
    is_xx_dtc = (
        combined_lib_df['Target Column'].astype(str).str.match(r'^[A-Za-z]{2,3}(DTC|STDTC)$') & 
        (combined_lib_df['Source Column'].astype(str).str.strip() == 'RECORDDATE')
    )
    is_xx_presp = combined_lib_df['Target Column'].astype(str).str.match(r'^[A-Za-z]{2}(PRESP)$')
    is_conditional_var = combined_lib_df['Target Column'].isin(conditional_vars_list) | is_xx_presp

    # A) TARGET ONLY RULES (Derived)
    for target in unique_target_tables:
        derived_df = combined_lib_df[(combined_lib_df['Source Table'] == 'DERIVED') & (combined_lib_df['Target Table'] == target)].copy()
        if not derived_df.empty: appended_rows_list.append(derived_df[output_headers])

    # B) TARGET/SOURCE CONDITIONAL MATCHES
    matched_mand_df = pd.merge(combined_lib_df[is_mand_list].copy(), active_table_pairs, on=['Source Table', 'Target Table'], how='inner')
    if not matched_mand_df.empty: appended_rows_list.append(matched_mand_df[output_headers])

    matched_spid_df = pd.merge(combined_lib_df[is_xx_spid].copy(), active_table_pairs, on=['Source Table', 'Target Table'], how='inner')
    if not matched_spid_df.empty: appended_rows_list.append(matched_spid_df[output_headers])

    matched_cat_scat_df = pd.merge(combined_lib_df[is_xx_cat_scat].copy(), active_table_pairs, on=['Source Table', 'Target Table'], how='inner')
    if not matched_cat_scat_df.empty: appended_rows_list.append(matched_cat_scat_df[output_headers])

    matched_dtc_df = pd.merge(combined_lib_df[is_xx_dtc].copy(), active_table_pairs, on=['Source Table', 'Target Table'], how='inner')
    if not matched_dtc_df.empty: appended_rows_list.append(matched_dtc_df[output_headers])

    matched_cond_vars_df = pd.merge(combined_lib_df[is_conditional_var].copy(), active_table_pairs, on=['Source Table', 'Target Table'], how='inner')
    if not matched_cond_vars_df.empty: appended_rows_list.append(matched_cond_vars_df[output_headers])
    
    matched_missing_src_df = pd.merge(combined_lib_df[is_missing_src_col].copy(), active_table_pairs, on=['Source Table', 'Target Table'], how='inner')
    if not matched_missing_src_df.empty: 
        appended_rows_list.append(matched_missing_src_df[output_headers])

    # Combine all
    final_df = pd.concat([core_mapped_df] + appended_rows_list, ignore_index=True) if appended_rows_list else core_mapped_df

    # ---------------------------------------------------------
    # STEP 9: DERIVED DATE/DAY DEPENDENCY RULE
    # ---------------------------------------------------------
    existing_targets = set(zip(final_df['Target Table'], final_df['Target Column']))
    rows_to_drop = []
    
    for idx, row in final_df.iterrows():
        if row['Source Table'] == 'DERIVED':
            tc = str(row['Target Column'])
            if re.match(r'^[A-Za-z]{2,3}(STDT|STDY|STEDY)$', tc):
                parent_tc = re.sub(r'(STDT|STDY|STEDY)$', 'STDTC', tc)
                if (row['Target Table'], parent_tc) not in existing_targets:
                    rows_to_drop.append(idx)
            elif re.match(r'^[A-Za-z]{2,3}(ENDT|ENDY|ENEDY)$', tc):
                parent_tc = re.sub(r'(ENDT|ENDY|ENEDY)$', 'ENDTC', tc)
                if (row['Target Table'], parent_tc) not in existing_targets:
                    rows_to_drop.append(idx)

    if rows_to_drop:
        final_df = final_df.drop(rows_to_drop).reset_index(drop=True)
        
    # ---------------------------------------------------------
    # STEP 10: RE-SEQUENCE TRANSPOSE COLUMN
    # ---------------------------------------------------------
    final_df = final_df.sort_values(by=['Target Table']).reset_index(drop=True)
    final_df = replace_trans(final_df)
    
   # ---------------------------------------------------------
    # OUTPUT CONSTRAINTS & SORTING
    # ---------------------------------------------------------
    
    final_df['Mapping Logic (pseudo code)'] = final_df['Mapping Logic (pseudo code)'].fillna('')

    # --- ADDED LOGIC FOR STUDYID ---
    study_id_mask = (
        (final_df['Source Table'] == 'SUBJECT_VISIT_BYDOMAIN') & 
        (final_df['Target Table'] == 'SV') & 
        (final_df['Target Column'] == 'STUDYID')
    )
    final_df.loc[study_id_mask, 'Mapping Logic (pseudo code)'] = f'"{studynum}"'

    # -------------------------------
    
    
    def check_if_really_unmatched(row):
        src_table = row['Source Table']
        src_col = str(row['Source Column'])
        
        table_logic = final_df[final_df['Source Table'] == src_table]['Mapping Logic (pseudo code)']
        if table_logic.empty:
            return True
            
        is_found_in_logic = table_logic.str.contains(src_col, regex=False).any()
        return not is_found_in_logic

    if not unmatched_df.empty:
        really_unmatched_mask = unmatched_df.apply(check_if_really_unmatched, axis=1)
        unmatched_df = unmatched_df[really_unmatched_mask]
    
    unmatched_df['is_unmatched'] = True
    final_df['is_unmatched'] = False
    
    final_df = pd.concat([unmatched_df, final_df], ignore_index=True)
    
    final_df = final_df[~final_df['Source Column'].astype(str).str.contains('LABEL_|LBL_', na=False, regex=True)]
    final_df = final_df.drop_duplicates()
    final_df['is_derived'] = final_df['Source Table'] == 'DERIVED'
    
    final_df = final_df.sort_values(
        by=['is_unmatched', 'Target Table', 'Source Table', 'Transpose', 'is_derived', 'Target Column', 'Source Column'],
        ascending=[False, True, True, True, True, True, True],
        na_position='last'
    )

    # ---------------------------------------------------------
    # EXPORT SUMMARY & FORMATTING
    # ---------------------------------------------------------
    total_output_rows = len(final_df)
    unmapped_output_rows = final_df['is_unmatched'].sum()
    mapped_output_rows = total_output_rows - unmapped_output_rows

    lib_total_rows = len(pd.read_excel(cdms_lib_file, sheet_name=cdms_lib_sheet))
    design_total_rows = len(design_df)
    design_total_forms = design_df['FormOID'].nunique() if 'FormOID' in design_df.columns else 0
    design_total_vars = len(design_df[['Derived_Source_Table', 'Derived_Source_Column']].drop_duplicates())
    
    unique_unmapped_vars = len(unmatched_df[['Source Table', 'Source Column']].drop_duplicates())
    unique_mapped_vars = design_total_vars - unique_unmapped_vars

    final_df = final_df[output_headers].reset_index(drop=True)
    mapped_pct = (unique_mapped_vars / (unique_mapped_vars + unmapped_output_rows) * 100) if (unique_mapped_vars + unmapped_output_rows) > 0 else 0.0

    print(f"\n------------------------------------------------")
    print(f"\nAs you sent the ALS for {studynum}, I have placed the CDMS which was generated by the Utility Tool, along with the help of reference studies. Is placed in the shared folder.")
    print(f"\nSummary")
    print(f"Design Spec Total Rows: {design_total_rows}")
    print(f"  -> Total rows to export: {total_output_rows}")
    print(f"  -> Total Unique Forms: {design_total_forms}")
    print(f"  -> Total Unique Variables: {design_total_vars}")
    print(f"  -> Variables Successfully Mapped: {unique_mapped_vars}")
    print(f"  -> Variables Unmapped Rows: {unmapped_output_rows}")
    print(f"  -> Percentage Mapped: {mapped_pct:.2f}%")
    print(f"Thank You! \n\n Best Regards, \n Tejasvin R")
    print(f"\n--- Final Export Summary ---")
    print(f"  -> Generated Mapped Rows: {mapped_output_rows}")
    print(f"  -> Variables Unmapped (Not Found): {unique_unmapped_vars}")
    print(f"CDMS Library Total Rows: {lib_total_rows}")
    print(f"Exporting to {output_file}...\n")

    track_study_metrics_in_excel(
        study_id=studynum,
        design_total_rows=design_total_rows,
        design_total_forms=design_total_forms,
        design_total_vars=design_total_vars,
        unique_mapped_vars=unique_mapped_vars,
        total_output_rows=total_output_rows,
        unmapped_output_rows=unmapped_output_rows,
        mapped_pct=mapped_pct,
        output_file=output_file,
    )
    
    # ---------------------------------------------------------
    # MULTI-SHEET EXPORT & FORMATTING LOGIC
    # ---------------------------------------------------------
    thin_side = Side(border_style="thin", color="000000")
    all_border = Border(top=thin_side, left=thin_side, right=thin_side, bottom=thin_side)
    grey_fill = PatternFill(start_color="808080", end_color="808080", fill_type="solid")
    black_bold_font = Font(bold=True, color="000000")

    with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
        approval_df.to_excel(writer, sheet_name='Approval', index=False, header=False)
        doc_history_df.to_excel(writer, sheet_name='Document History', index=False, header=False)
        final_df.to_excel(writer, sheet_name='Database Mapping', index=False)
        
        workbook = writer.book
        
        try:
            source_wb = openpyxl.load_workbook(cdms_lib_file, data_only=True)
            def copy_sheet_styles(sheet_name):
                if sheet_name not in source_wb.sheetnames: return
                src_ws = source_wb[sheet_name]
                tgt_ws = workbook[sheet_name]
                
                for col_letter, col_dim in src_ws.column_dimensions.items():
                    tgt_ws.column_dimensions[col_letter].width = col_dim.width

                for row_num, row_dim in src_ws.row_dimensions.items():
                    tgt_ws.row_dimensions[row_num].height = row_dim.height

                for merged_range in src_ws.merged_cells.ranges:
                    tgt_ws.merge_cells(merged_range.coord)

                for row in src_ws.iter_rows():
                    for cell in row:
                        tgt_cell = tgt_ws[cell.coordinate]
                        if cell.has_style:
                            tgt_cell.font = copy.copy(cell.font)
                            tgt_cell.border = copy.copy(cell.border)
                            tgt_cell.fill = copy.copy(cell.fill)
                            tgt_cell.number_format = copy.copy(cell.number_format)
                            tgt_cell.alignment = copy.copy(cell.alignment)
                            
            copy_sheet_styles('Approval')
            copy_sheet_styles('Document History')
            
        except Exception as e:
            print(f"Warning: Could not strictly copy styles from source library. Reason: {e}")

        ws_db = workbook['Database Mapping']
        for row in ws_db['A1:K1']:
            for cell in row:
                cell.font = black_bold_font
                cell.fill = grey_fill
                
        max_row = ws_db.max_row
        for row in ws_db[f'A1:K{max_row}']:
            for cell in row:
                cell.border = all_border

        ws_db.column_dimensions['A'].width = 15
        ws_db.column_dimensions['B'].width = 20
        ws_db.column_dimensions['C'].width = 20
        ws_db.column_dimensions['D'].width = 10
        ws_db.column_dimensions['E'].width = 20
        ws_db.column_dimensions['F'].width = 10
        ws_db.column_dimensions['I'].width = 30
        ws_db.column_dimensions['K'].width = 25
            
    print("Execution complete!")

if __name__ == "__main__":
    generate_cdms_mapping()